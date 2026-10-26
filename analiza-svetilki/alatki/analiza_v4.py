"""Анализа на стари светилки – верзија 4 (енергетски биланс по населено место и месец).

Претпоставки (договорени со корисникот):
  - загуби по броилото 10 %;
  - LED од ЈПП: бројот и моќноста од списокот се точни, светат 100 % од времето на часовникот;
  - стари светилки: 10 W, 80 % активни (20 % неисправни);
  - астрономски часовник: од зајдисонце до изгрејсонце (проверка со ±15 минути);
  - мерните места се правилно распоредени по места; градот Кичево е една целина;
  - селата без LED во списокот имаат само стари светилки;
  - украсување 15.12–15.01 на броилата од vlez/ukrasi.csv (проценето како вишок
    во декември и јануари над вообичаената потрошувачка на тие броила);
  - LV2 броилата не влегуваат.

Пресметка за секое место и месец:
  корисни kWh  = (kWh − украсување) × 0,9
  LED kWh      = LED kW × часови на работа
  стари        = (корисни − LED kWh) ÷ (10 W × часови) ÷ 0,8

Годишна проценка (главна): од енергијата за 12 месеци (не зависи од поместени читања).
Проверка 1: медијана на месечните проценки.
Проверка 2 (независна): сезонска поделба на месечните kWh на ноќен дел и дел
што не зависи од должината на ноќта; моќноста од ноќниот дел дава број стари.

Употреба:  python3 analiza-svetilki/alatki/analiza_v4.py
Излез:     izlez/v4_*.csv и izlez/Анализа_светилки_v4.xlsx
"""
import calendar
import csv
import statistics
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import analiza as A

LOSS, W_OLD, ACTIVE = 0.10, 10.0, 0.8
DECO_MONTHS = ("2025-12", "2026-01")
CABINETS = {"Осломеј": 35, "Зајас": 49, "Вранештица": 24, "Другово": 39, "Кичево": 87}
OK, WARN, BAD = (PatternFill("solid", fgColor=c) for c in ("E3F1E3", "FFF4CC", "FDE2E1"))


def fit(v, months, H, D):
    """kWh ≈ a × ноќни часови + b × денови; a = моќност (kW), b ÷ a = часови на ден над ноќта."""
    ms = [m for m in months if m in v]
    x1, x2, y = [H[m] for m in ms], [D[m] for m in ms], [v[m] for m in ms]
    if sum(1 for e in y if e > 0) < 4:
        return None, None
    s11, s22, s12 = sum(p * p for p in x1), sum(q * q for q in x2), sum(p * q for p, q in zip(x1, x2))
    t1, t2 = sum(p * e for p, e in zip(x1, y)), sum(q * e for q, e in zip(x2, y))
    det = s11 * s22 - s12 * s12
    a, b = (t1 * s22 - t2 * s12) / det, (s11 * t2 - s12 * t1) / det
    if a < 0:
        return 0.0, t2 / s22
    if b < 0:
        return t1 / s11, 0.0
    return a, b


def load():
    places, lamps, mapping = A.read("naseleni.csv"), A.read("svetilki.csv"), A.read("mapiranje.csv")
    deco_meters = {r["meter"] for r in A.read("ukrasi.csv")}
    group = {p["unit"]: p["group"] for p in places}
    members = defaultdict(list)
    for p in places:
        members[p["unit"]].append(p["settlement"])
    led_n, led_kw = Counter(), Counter()
    for l in lamps:
        led_n[l["unit"]] += int(l["count"])
        led_kw[l["unit"]] += int(l["count"]) * float(l["watt"]) / 1000
    unit_of = {r["meter"]: r["unit"] for r in mapping if r["unit"]}
    per_meter = defaultdict(dict)
    total_invoice = 0.0
    for r in A.read("stavki.csv"):
        if "LV1" not in r["category"]:
            continue
        total_invoice += float(r["kwh"])
        m = A.month_key(r["month"])
        per_meter[r["meter"]][m] = per_meter[r["meter"]].get(m, 0) + float(r["kwh"])
    months = sorted({m for v in per_meter.values() for m in v})
    H = A.night_hours(months)
    D = {m: calendar.monthrange(int(m[:4]), int(m[5:]))[1] for m in months}
    return places, group, members, led_n, led_kw, unit_of, per_meter, deco_meters, months, H, D, total_invoice


def analyse():
    places, group, members, led_n, led_kw, unit_of, per_meter, deco_meters, months, H, D, total_invoice = load()
    no_deco = [m for m in months if m not in DECO_MONTHS]
    E = defaultdict(lambda: defaultdict(float))
    deco = defaultdict(lambda: defaultdict(float))
    meters = defaultdict(set)
    deco_rows = []
    for meter, v in per_meter.items():
        u = unit_of[meter]
        meters[u].add(meter)
        a, b = fit(v, no_deco, H, D) if meter in deco_meters else (None, None)
        for m in months:
            d = max(0.0, v.get(m, 0) - a * H[m] - b * D[m]) if a is not None and m in DECO_MONTHS else 0.0
            E[u][m] += v.get(m, 0) - d
            deco[u][m] += d
        if meter in deco_meters:
            deco_rows.append({"meter": meter, "kwh": round(sum(v.values())), "dek": round(v.get("2025-12", 0)), "jan": round(v.get("2026-01", 0)),
                              "ocekuvano_dek": round(a * H["2025-12"] + b * D["2025-12"]), "ocekuvano_jan": round(a * H["2026-01"] + b * D["2026-01"]),
                              "ukras": round(sum(max(0.0, v.get(m, 0) - a * H[m] - b * D[m]) for m in DECO_MONTHS))})

    def old(useful, kw, h, w=W_OLD, act=ACTIVE):
        return (useful - kw * h) / (w / 1000 * h) / act

    monthly, annual = [], []
    for u in sorted(set(group), key=lambda x: (A.GROUPS.index(group[x]), -sum(E[x].values()))):
        kw = led_kw[u]
        mvals = []
        for m in months:
            useful = E[u][m] * (1 - LOSS)
            n = old(useful, kw, H[m]) if meters[u] else None
            monthly.append({"opstina": group[u], "mesto": u, "mesec": m, "nokni_casovi": round(H[m], 1), "kwh": round(E[u][m] + deco[u][m], 1),
                            "ukras": round(deco[u][m], 1), "korisni_kwh": round(useful, 1), "led": led_n[u], "led_kwh": round(kw * H[m], 1),
                            "ostatok_kwh": round(useful - kw * H[m], 1), "stari_aktivni": None if n is None else round(n * ACTIVE),
                            "stari_postaveni": None if n is None else round(n)})
            if n is not None and E[u][m] > 0:
                mvals.append(n)
        tot = sum(E[u].values())
        useful_y = tot * (1 - LOSS)
        hy = sum(H.values())
        if not meters[u]:
            annual.append({"opstina": group[u], "mesto": u, "naseleni_mesta": ", ".join(members[u]), "merni_mesta": 0, "led": led_n[u], "led_kw": round(kw, 2),
                           "kwh": 0, "ukras": 0, "korisni_kwh": 0, "led_kwh": round(kw * hy), "stari_godisno": None, "stari_medijana": None,
                           "stari_sezonski": None, "dopolnitelni_casovi": None, "stari_proceneti": 0, "razlika": None, "status": "Нема мерно место"})
            continue
        n_year = old(useful_y, kw, hy)
        med = statistics.median(mvals) if mvals else None
        a, b = fit({m: E[u][m] * (1 - LOSS) for m in months}, months, H, D)
        n_season = (a - kw) / (W_OLD / 1000) / ACTIVE if a is not None else None
        extra = b / a if a else None
        spread = abs(n_year - med) / max(abs(n_year), 20) if med is not None else 0
        if n_year < 0:
            status = "LED трошат повеќе од измереното"
        elif not led_n[u]:
            status = "Само стари светилки"
        elif spread > 0.3:
            status = "Нерамномерни месеци – проверете"
        else:
            status = "LED и стари светилки"
        annual.append({"opstina": group[u], "mesto": u, "naseleni_mesta": ", ".join(members[u]), "merni_mesta": len(meters[u]), "led": led_n[u], "led_kw": round(kw, 2),
                       "kwh": round(tot + sum(deco[u].values())), "ukras": round(sum(deco[u].values())), "korisni_kwh": round(useful_y), "led_kwh": round(kw * hy),
                       "stari_godisno": round(n_year), "stari_medijana": None if med is None else round(med),
                       "stari_sezonski": None if n_season is None else round(n_season), "dopolnitelni_casovi": None if extra is None else round(extra, 1),
                       "stari_proceneti": max(0, round(n_year)), "razlika": round(spread, 2), "status": status})

    groups = []
    for g in A.GROUPS + ["Вкупно"]:
        rs = [r for r in annual if g == "Вкупно" or r["opstina"] == g]
        net = sum(r["stari_godisno"] or 0 for r in rs)
        groups.append({"opstina": g, "led": sum(r["led"] for r in rs), "merni_mesta": sum(r["merni_mesta"] for r in rs),
                       "ormari": sum(CABINETS.values()) if g == "Вкупно" else CABINETS[g],
                       "kwh": sum(r["kwh"] for r in rs), "ukras": sum(r["ukras"] for r in rs), "korisni_kwh": sum(r["korisni_kwh"] for r in rs),
                       "led_kwh": sum(r["led_kwh"] for r in rs), "stari_neto": round(net), "stari_zbir_pozitivni": sum(r["stari_proceneti"] for r in rs),
                       "stari_sezonski": round(sum(r["stari_sezonski"] or 0 for r in rs)),
                       "stari_mesecno": {m: round(sum(x["stari_postaveni"] or 0 for x in monthly if x["mesec"] == m and (g == "Вкупно" or x["opstina"] == g))) for m in months}})

    def scenario(shift=0.0, loss=LOSS, ledf=1.0, w=W_OLD, act=ACTIVE):
        out = Counter()
        hy = sum(H.values()) + shift * sum(D.values())
        for u in group:
            if meters[u]:
                out[group[u]] += old(sum(E[u].values()) * (1 - loss), led_kw[u] * ledf, hy, w, act)
        out["Вкупно"] = sum(out[g] for g in A.GROUPS)
        city = old(sum(E["Кичево (град)"].values()) * (1 - loss), led_kw["Кичево (град)"] * ledf, hy, w, act)
        return out, city

    scen_defs = [
        ("Основно: зајдисонце–изгрејсонце", {}),
        ("Вклучување и исклучување 15 мин порано/подоцна (−0,5 ч/ден)", {"shift": -0.5}),
        ("Вклучување и исклучување 15 мин подолго (+0,5 ч/ден)", {"shift": +0.5}),
        ("Само вклучување 15 мин подоцна (−0,25 ч/ден)", {"shift": -0.25}),
        ("Само вклучување 15 мин порано (+0,25 ч/ден)", {"shift": +0.25}),
        ("Загуби 7 %", {"loss": 0.07}), ("Загуби 13 %", {"loss": 0.13}),
        ("LED трошат 5 % помалку од номиналното", {"ledf": 0.95}), ("LED трошат 5 % повеќе од номиналното", {"ledf": 1.05}),
        ("Стари 8 W", {"w": 8}), ("Стари 12 W", {"w": 12}), ("Активни стари 70 %", {"act": 0.7}), ("Активни стари 90 %", {"act": 0.9}),
        ("Комбинација: најмалку стари", {"shift": 0.5, "loss": 0.13, "ledf": 1.05, "w": 12, "act": 0.9}),
        ("Комбинација: најмногу стари", {"shift": -0.5, "loss": 0.07, "ledf": 0.95, "w": 8, "act": 0.7}),
    ]
    scenarios = []
    for name, kw in scen_defs:
        out, city = scenario(**kw)
        scenarios.append({"scenario": name, **{g: round(out[g]) for g in A.GROUPS + ["Вкупно"]}, "grad": round(city)})

    base = scenarios[0]["Вкупно"]
    singles = [s["Вкупно"] - base for s in scenarios[1:13]]
    pairs = [max(abs(singles[i]), abs(singles[i + 1])) for i in (0, 4, 6, 8, 10)]
    rss = round(sum(p * p for p in pairs) ** 0.5)

    checks = {
        "kwh_faktura": round(total_invoice), "kwh_vo_analiza": round(sum(r["kwh"] for r in annual)),
        "led_spisok": sum(led_n.values()), "led_vo_analiza": sum(r["led"] for r in annual),
        "merni_mesta": len(per_meter), "merni_mesta_rasporedeni": sum(len(v) for v in meters.values()),
        "nesigurnost_vkupno": rss,
    }
    return months, H, monthly, annual, groups, scenarios, deco_rows, checks


def write(months, H, monthly, annual, groups, scenarios, deco_rows, checks):
    for name, rows in (("v4_po_mesto_mesec.csv", monthly), ("v4_po_mesto.csv", annual), ("v4_scenarija.csv", scenarios)):
        with open(A.IZLEZ / name, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader(); w.writerows(rows)
    gflat = [{k: v for k, v in g.items() if k != "stari_mesecno"} | {f"stari_{m}": g["stari_mesecno"][m] for m in months} for g in groups]
    with open(A.IZLEZ / "v4_po_opstina.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(gflat[0].keys()))
        w.writeheader(); w.writerows(gflat)

    bold, head = Font(bold=True), PatternFill("solid", fgColor="E8EEF2")
    wb = Workbook()

    def sheet(title, header, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
        for i, wd in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wd
        ws.freeze_panes = "C2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.0"
                elif isinstance(c.value, int):
                    c.number_format = "#,##0"
        return ws

    ws = wb.active
    ws.title = "Пристап и точност"
    tot = next(g for g in groups if g["opstina"] == "Вкупно")
    city = next(r for r in annual if r["mesto"] == "Кичево (град)")
    lines = [
        ("Анализа на стари светилки – верзија 4 (енергетски биланс по населено место и месец)", True),
        ("", False),
        ("Претпоставки", True),
        ("Загуби по броилото 10 %. LED од ЈПП: број и моќност од списокот, светат 100 % од времето на часовникот.", False),
        ("Стари светилки: 10 W, 80 % активни. Астрономски часовник: зајдисонце–изгрејсонце (проверка ±15 мин).", False),
        ("Мерните места се правилно распоредени; градот Кичево е една целина; селата без LED имаат само стари светилки.", False),
        ("Украсување 15.12–15.01 на 12 броила во градот, проценето како вишок во декември и јануари. LV2 броилата не влегуваат.", False),
        ("", False),
        ("Пресметка", True),
        ("За секое место и месец: стари = ((kWh − украсување) × 0,9 − LED kW × часови) ÷ (10 W × часови) ÷ 0,8.", False),
        ("Годишна проценка (главна): од енергијата за 12 месеци – не зависи од поместени или проценети читања.", False),
        ("Проверка 1: медијана на месечните проценки. Проверка 2: сезонска поделба (моќност од ноќниот дел на потрошувачката).", False),
        ("По поранешна општина: нето збир (местата каде LED трошат повеќе од измереното се одземаат).", False),
        ("", False),
        ("Резултат", True),
        (f"Вкупно стари светилки (нето): {tot['stari_neto']:,}; збир само на позитивните места: {tot['stari_zbir_pozitivni']:,}; сезонска проверка: {tot['stari_sezonski']:,}.".replace(",", "."), False),
        (f"Градот Кичево: {city['stari_godisno']:,} стари светилки при зајдисонце–изгрејсонце – измерената енергија е колку што трошат само LED.".replace(",", "."), False),
        ("", False),
        ("Ниво на точност", True),
        (f"Бројот на стари светилки е остаток од разлика на два големи броја (корисни kWh и LED kWh), па е многу осетлив на претпоставките.", False),
        (f"Проценета несигурност на вкупниот број (комбинација на часовник ±15 мин, загуби ±3 %, LED ±5 %, 8–12 W, 70–90 % активни): ± {checks['nesigurnost_vkupno']:,} светилки.".replace(",", "."), False),
        ("Места каде LED трошат речиси сè (градот, Гарани и Стрелци, Грешница, Србица): несигурноста е поголема од самиот резултат.", False),
        ("Места без LED (Раштани, Душегубица, Јаворец, села на Вранештица): ±30–40 %, зависи само од моќноста и уделот на активни стари светилки.", False),
        ("", False),
        ("Двојна проверка на пресметката", True),
        (f"kWh во фактурите (LV1.1): {checks['kwh_faktura']:,}; kWh распоредени по места: {checks['kwh_vo_analiza']:,}.".replace(",", "."), False),
        (f"LED во списокот: {checks['led_spisok']:,}; LED во анализата: {checks['led_vo_analiza']:,}.".replace(",", "."), False),
        (f"Мерни места: {checks['merni_mesta']}; распоредени: {checks['merni_mesta_rasporedeni']}.", False),
        ("Ормари (список на ЈПП) наспроти мерни места по поранешна општина: видете лист „По општина“.", False),
    ]
    for text, is_head in lines:
        ws.append([text])
        if is_head:
            ws[ws.max_row][0].font = Font(bold=True, size=13 if ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 140

    mh = [A.SHORT[m[5:]] + " " + m[2:4] for m in months]
    sheet("По општина", ["Поранешна општина", "LED", "Ормари", "Мерни места", "kWh", "Украсување", "Корисни kWh", "LED kWh", "Стари (нето)", "Стари (збир позитивни)",
                         "Стари (сезонска проверка)"] + ["Стари " + h for h in mh],
          [[g["opstina"], g["led"], g["ormari"], g["merni_mesta"], g["kwh"], g["ukras"], g["korisni_kwh"], g["led_kwh"], g["stari_neto"], g["stari_zbir_pozitivni"],
            g["stari_sezonski"]] + [g["stari_mesecno"][m] for m in months] for g in groups], [18, 8, 8, 8, 11, 10, 11, 11, 10, 11, 11] + [8] * len(months))
    ws2 = sheet("По место", ["Поранешна општина", "Населено место", "Населени места", "Мерни места", "LED", "LED kW", "kWh", "Украсување", "Корисни kWh", "LED kWh",
                             "Стари (годишно)", "Стари (медијана месеци)", "Стари (сезонска проверка)", "Часови/ден над ноќта (проверка)", "Разлика год./медијана", "Статус"],
                [[r["opstina"], r["mesto"], r["naseleni_mesta"], r["merni_mesta"], r["led"], r["led_kw"], r["kwh"], r["ukras"], r["korisni_kwh"], r["led_kwh"],
                  r["stari_godisno"], r["stari_medijana"], r["stari_sezonski"], r["dopolnitelni_casovi"], r["razlika"], r["status"]] for r in annual],
                [12, 22, 30, 8, 7, 8, 10, 9, 10, 10, 10, 10, 10, 10, 9, 30])
    for row in ws2.iter_rows(min_row=2):
        st = row[15].value
        fill = OK if st in ("LED и стари светилки", "Само стари светилки") else WARN if st and st.startswith("Нерамномерни") else BAD if st and st.startswith("LED трошат") else None
        if fill:
            row[15].fill = fill
    by = {(r["mesto"], r["mesec"]): r for r in monthly}
    sheet("Место × месец", ["Поранешна општина", "Населено место", "LED"] + ["Стари " + h for h in mh] + ["Годишно", "Медијана"],
          [[r["opstina"], r["mesto"], r["led"]] + [by[r["mesto"], m]["stari_postaveni"] for m in months] + [r["stari_godisno"], r["stari_medijana"]] for r in annual],
          [12, 22, 7] + [8] * len(months) + [9, 9])
    sheet("Детално по месец", ["Поранешна општина", "Населено место", "Месец", "Ноќни часови", "kWh", "Украсување", "Корисни kWh", "LED", "LED kWh", "Остаток kWh",
                               "Активни стари", "Поставени стари"],
          [[r["opstina"], r["mesto"], r["mesec"], r["nokni_casovi"], r["kwh"], r["ukras"], r["korisni_kwh"], r["led"], r["led_kwh"], r["ostatok_kwh"],
            r["stari_aktivni"], r["stari_postaveni"]] for r in monthly], [12, 22, 9, 9, 10, 9, 10, 7, 10, 10, 9, 9])
    sheet("Сценарија (точност)", ["Сценарио"] + A.GROUPS + ["Вкупно", "од тоа градот"],
          [[s["scenario"]] + [s[g] for g in A.GROUPS + ["Вкупно"]] + [s["grad"]] for s in scenarios], [58, 10, 10, 10, 11, 10, 10, 12])
    deco_names = {r["meter"]: r["zabeleska"] for r in A.read("ukrasi.csv")}
    sheet("Украсување", ["Мерно место", "Локација", "kWh 12 мес.", "Декември", "Очекувано дек.", "Јануари", "Очекувано јан.", "Украсување kWh"],
          [[d["meter"], deco_names[d["meter"]], d["kwh"], d["dek"], d["ocekuvano_dek"], d["jan"], d["ocekuvano_jan"], d["ukras"]] for d in deco_rows],
          [17, 18, 11, 10, 12, 10, 12, 12])
    wb.save(A.IZLEZ / "Анализа_светилки_v4.xlsx")


def main():
    months, H, monthly, annual, groups, scenarios, deco_rows, checks = analyse()
    write(months, H, monthly, annual, groups, scenarios, deco_rows, checks)
    for g in groups:
        print(f"{g['opstina']:11} LED {g['led']:5} ормари {g['ormari']:3} мерни {g['merni_mesta']:3} | стари нето {g['stari_neto']:6} позитивни {g['stari_zbir_pozitivni']:6} сезонски {g['stari_sezonski']:6}")
    print(checks)
    for s in scenarios:
        print(f"{s['scenario'][:58]:58} вкупно {s['Вкупно']:6} град {s['grad']:5}")


if __name__ == "__main__":
    main()
