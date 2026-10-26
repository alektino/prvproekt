"""Проба: дали разликата се должи на LED светилки различни од пописот?

Хипотеза на корисникот: LED се вградувани без надзор, „што имало“. Во селата
можеби се вградени посилни, во градот можеби послаби од пописот на ЈПП.

Бројот на стари светилки се фиксира (избројан во контролните села, ваша проценка
по поранешна општина, 500–600 во градот), а се пресметува колкава LED моќност
треба да е вградена за да се објасни измереното:

  потребна LED kW = корисни kWh ÷ часови − стари × 10 W × 0,8
  корисни kWh     = (kWh − украсување) × 0,9;  часови = зајдисонце–изгрејсонце

Моќностите во пописот се 12,1 / 20,9 / 25,9 / 32,1 / 64,5 W. Ако потребната
просечна моќност е над 64,5 W, хипотезата сама не ја објаснува потрошувачката.

Употреба (по analiza_v4.py):  python3 analiza-svetilki/alatki/proba_led.py
Излез: izlez/proba_led.csv и izlez/Проба_LED.xlsx
"""
import csv
import statistics

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import analiza as A
import analiza_v4 as V

W_EFF = V.W_OLD * V.ACTIVE / 1000          # kW по поставена стара светилка
W_MAX, W_EAE = 64.5, 25.9
COUNTED = {"Староец": (29, 26), "Видрани": (7, 23)}   # избројани LED, стари
CITY = "Кичево (град)"
CITY_OLD = (500, 600)


def main():
    months, H, monthly, annual, groups, scenarios, deco_rows, checks = V.analyse()
    hy = sum(H.values())
    rows = {r["mesto"]: r for r in annual if r["merni_mesta"]}
    share = {r["opstina"]: float(r["udel_stari"]) for r in A.read("procenki.csv")}
    out = []

    def row(kind, name, led_list, led_kw_list, n_old, power, note=""):
        need = power - n_old * W_EFF
        avg = need * 1000 / led_list if led_list else None
        swaps = (need - led_kw_list) * 1000 / (W_MAX - W_EAE)
        if avg is None:
            verdict = ""
        elif avg > W_MAX:
            verdict = "Не е можно ни со најсилните LED"
        elif need >= led_kw_list:
            verdict = "Можно: посилни LED од пописот"
        else:
            verdict = "Можно: послаби LED од пописот"
        out.append({"vid": kind, "mesto": name, "led": led_list, "led_kw_popis": round(led_kw_list, 2), "stari": round(n_old),
                    "izmereno_kw": round(power, 2), "potrebno_led_kw": round(need, 2), "odnos": round(need / led_kw_list, 2) if led_kw_list else None,
                    "prosek_w_popis": round(led_kw_list * 1000 / led_list, 1) if led_list else None, "prosek_w_potrebno": None if avg is None else round(avg, 1),
                    "zameni_25_9_so_64_5": round(swaps), "zaklucok": verdict, "zabeleska": note})

    # 1. Контролни села: медијана на месечната моќност (неправилните читања не влијаат).
    lamps = A.read("svetilki.csv")
    unit_of = {r["meter"]: r["unit"] for r in A.read("mapiranje.csv") if r["unit"]}
    per = {}
    for r in A.read("stavki.csv"):
        if "LV1" in r["category"] and unit_of.get(r["meter"]) in COUNTED:
            k = (unit_of[r["meter"]], A.month_key(r["month"]))
            per[k] = per.get(k, 0) + float(r["kwh"])
    for u, (n_led, n_old) in COUNTED.items():
        power = statistics.median(per[u, m] * (1 - V.LOSS) / H[m] for m in months)
        listed = sum(int(l["count"]) for l in lamps if l["unit"] == u)
        row("Избројано село", u, n_led, n_led * W_EAE / 1000, n_old, power,
            f"Попис на ЈПП: {listed} LED по {W_EAE} W; моќност = медијана на месечната kW")

    # 2. Град: 500–600 стари.
    c = rows[CITY]
    for n in CITY_OLD:
        row("Град", f"{CITY}, {n} стари", c["led"], c["led_kw"], n, c["korisni_kwh"] / hy)

    # 3. Поранешни општини: стари по ваша проценка (удел од сите светилки).
    for g in A.GROUPS:
        rs = [r for r in annual if r["opstina"] == g and r["merni_mesta"]]
        led, led_kw = sum(r["led"] for r in rs), sum(r["led_kw"] for r in rs)
        n_old = led * share[g] / (1 - share[g])
        row("Поранешна општина", g, led, led_kw, n_old, sum(r["korisni_kwh"] for r in rs) / hy, f"Стари = {share[g]:.0%} од сите светилки")
    gs = [r for r in out if r["vid"] == "Поранешна општина"]
    row("Вкупно", "Сите (ваша проценка по општина)", sum(r["led"] for r in gs), sum(r["led_kw_popis"] for r in gs), sum(r["stari"] for r in gs),
        sum(r["izmereno_kw"] for r in gs))
    total = out[-1]
    row("Вкупно", "Сите (1.500 стари)", total["led"], total["led_kw_popis"], 1500, total["izmereno_kw"])

    with open(A.IZLEZ / "proba_led.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
        w.writeheader(); w.writerows(out)
    write_xlsx(out, scenarios[0]["Вкупно"])
    for r in out:
        print(f"{r['vid']:18} {r['mesto']:34} LED {r['led']:5} стари {r['stari']:5}  попис {r['led_kw_popis']:7.2f} kW → потребно {r['potrebno_led_kw']:7.2f} kW "
              f"(×{r['odnos']}, {r['prosek_w_potrebno']} W/св.)  {r['zaklucok']}")


def write_xlsx(out, v4_total):
    wb = Workbook()
    ws = wb.active
    ws.title = "Проба – LED"
    bold, head = Font(bold=True), PatternFill("solid", fgColor="E8EEF2")
    for line in [
        "Проба: колкава LED моќност треба да е вградена ако бројот на стари светилки е познат?",
        "Потребна LED kW = корисни kWh ÷ часови − стари × 10 W × 0,8. Пописот има LED од 12,1 / 20,9 / 25,9 / 32,1 / 64,5 W.",
        "„Замени 25,9 → 64,5“: колку светилки од 25,9 W треба да се всушност 64,5 W (негативно: обратно – послаби од пописот).",
        f"Важно: ако LED само се преместени меѓу местата (ист вкупен број по моќност), вкупниот број стари останува {v4_total:,} – се менува само распределбата.".replace(",", "."),
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    cols = ["Вид", "Место", "LED", "LED kW (попис)", "Стари (фиксно)", "Измерено kW", "Потребно LED kW", "Потребно ÷ попис", "W по LED (попис)",
            "W по LED (потребно)", "Замени 25,9 → 64,5", "Заклучок", "Забелешка"]
    ws.append(cols)
    for c in ws[ws.max_row]:
        c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="top")
    fills = {"Не": V.BAD, "Можно: посилни": V.WARN, "Можно: послаби": V.OK}
    for r in out:
        ws.append(list(r.values()))
        for key, fill in fills.items():
            if r["zaklucok"].startswith(key):
                ws.cell(ws.max_row, 12).fill = fill
    for col, w in zip("ABCDEFGHIJKLM", [18, 32, 8, 11, 10, 10, 11, 10, 10, 11, 11, 30, 44]):
        ws.column_dimensions[col].width = w
    wb.save(A.IZLEZ / "Проба_LED.xlsx")


if __name__ == "__main__":
    main()
