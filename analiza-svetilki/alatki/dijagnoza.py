"""Дијагноза на мерните места: каде и зошто сметките се поголеми од очекуваното.

Фактурите ја делат енергијата на висока тарифа (ВТ, 07–22 ч, понеделник–сабота)
и ниска тарифа (НТ, останатото време и цела недела). Светилка под астрономски
часовник троши во ВТ само во темниот дел пред 22 ч и по 07 ч (зиме), па уделот
на ВТ е однапред познат (проверено на Староец и на броилото 15737 во Видрани).
Секој товар што работи и преку ден (кражба/приклучок, светилки што светат и дење,
грејач во ормар, пумпа ...) го зголемува уделот на ВТ.

За секое броило, од 12 месеци × 2 тарифи (24 равенки) се пресметуваат:
  L – моќност што работи само ноќе, по часовник (kW): светилки;
  X – моќност што работи 24 часа (kW): сè друго.
  ВТ = L × темни часови во ВТ + X × сите часови во ВТ;  НТ = исто за НТ.
Месеци што отстапуваат над 35 % од пресметаното се означени како неправилни
читања (проценето читање, дополнителна наплата, корекција).

Очекувано по место = (LED од пописот + стари по вашата проценка за поранешната
општина × 10 W × 0,8) × ноќни часови ÷ 0,9.

Употреба:  python3 analiza-svetilki/alatki/dijagnoza.py
Излез:     izlez/dijagnoza_broila.csv, izlez/dijagnoza_mesta.csv, izlez/Дијагноза_сметки.xlsx
"""
import calendar
import csv
import datetime as dt
import math
import statistics
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import analiza as A
import analiza_v4 as V

HT = (7, 22)
SPIKE = 0.35
W_EFF = V.W_OLD * V.ACTIVE / 1000
OK, WARN, BAD = V.OK, V.WARN, V.BAD


def sun(d, lat=41.5127, lon=20.9589):
    """Изгрејсонце и зајдисонце за Кичево (NOAA), локално време."""
    rad, y, m = math.pi / 180, d.year, d.month
    last = lambda mm: (dt.date(y, mm + 1, 1) - dt.timedelta(days=1))
    lsun = lambda mm: last(mm).day - (last(mm).weekday() + 1) % 7
    g = 2 * math.pi / (366 if y % 4 == 0 else 365) * (d.timetuple().tm_yday - 1)
    eq = 229.18 * (0.000075 + 0.001868 * math.cos(g) - 0.032077 * math.sin(g) - 0.014615 * math.cos(2 * g) - 0.040849 * math.sin(2 * g))
    dec = 0.006918 - 0.399912 * math.cos(g) + 0.070257 * math.sin(g) - 0.006758 * math.cos(2 * g) + 0.000907 * math.sin(2 * g) - 0.002697 * math.cos(3 * g) + 0.00148 * math.sin(3 * g)
    ha = math.acos(math.cos(90.833 * rad) / (math.cos(lat * rad) * math.cos(dec)) - math.tan(lat * rad) * math.tan(dec)) / rad
    dst = 3 < m < 10 or (m == 3 and d.day >= lsun(3)) or (m == 10 and d.day < lsun(10))
    tz = 2 if dst else 1
    return (720 - 4 * (lon + ha) - eq) / 60 + tz, (720 - 4 * (lon - ha) - eq) / 60 + tz


def tariff_hours(k):
    """Темни часови и сите часови во ВТ и НТ за месецот."""
    y, m = map(int, k.split("-"))
    o = Counter()
    for day in range(1, calendar.monthrange(y, m)[1] + 1):
        d = dt.date(y, m, day)
        rise, set_ = sun(d)
        high = d.weekday() != 6
        dark = rise + 24 - set_
        dvt = max(0, rise - HT[0]) + max(0, HT[1] - set_) if high else 0
        o["dvt"] += dvt; o["dnt"] += dark - dvt
        o["tvt"] += HT[1] - HT[0] if high else 0; o["tnt"] += 24 - (HT[1] - HT[0] if high else 0)
    return o


def fit(rows, T):
    """Најмали квадрати за L (ноќно) и X (24 ч), обата ≥ 0."""
    eq = [(T[m]["dvt"], T[m]["tvt"], vt) for m, vt, nt in rows] + [(T[m]["dnt"], T[m]["tnt"], nt) for m, vt, nt in rows]
    s11 = sum(a * a for a, b, y in eq); s22 = sum(b * b for a, b, y in eq); s12 = sum(a * b for a, b, y in eq)
    t1 = sum(a * y for a, b, y in eq); t2 = sum(b * y for a, b, y in eq)
    det = s11 * s22 - s12 * s12
    L, X = (t1 * s22 - t2 * s12) / det, (s11 * t2 - s12 * t1) / det
    if X < 0:
        L, X = t1 / s11, 0.0
    if L < 0:
        L, X = 0.0, t2 / s22
    return L, X


def analyse():
    mapping = {r["meter"]: r for r in A.read("mapiranje.csv")}
    deco = {r["meter"] for r in A.read("ukrasi.csv")}
    raw = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0, 0.0]))
    for r in A.read("stavki.csv"):
        if "LV1" in r["category"]:
            v = raw[r["meter"]][A.month_key(r["month"])]
            v[0] += float(r["kwh_avt"]); v[1] += float(r["kwh_ant"]); v[2] += float(r["billed"])
    months = sorted({m for v in raw.values() for m in v})
    T = {m: tariff_hours(m) for m in months}
    hy = sum(T[m]["dvt"] + T[m]["dnt"] for m in months)

    meters = []
    for meter, v in raw.items():
        rows = [(m, v[m][0], v[m][1]) for m in months if m in v and v[m][0] + v[m][1] > 0]
        use = [r for r in rows if not (meter in deco and r[0] in V.DECO_MONTHS)]
        if len(use) < 4:
            continue
        # Најлошиот месец се отстранува додека сите не се во рамки од SPIKE (најмногу 4 месеци),
        # за неколку неправилни читања да не ја расипат пресметката за останатите.
        clean, spikes = list(use), []
        while True:
            L, X = fit(clean, T)
            pred = lambda m, L=L, X=X: L * (T[m]["dvt"] + T[m]["dnt"]) + X * (T[m]["tvt"] + T[m]["tnt"])
            dev = [(abs(vt + nt - pred(m)) / pred(m) if pred(m) > 0 else 0, m) for m, vt, nt in clean]
            worst, wm = max(dev)
            if worst <= SPIKE or len(spikes) == 4 or len(clean) <= 6:
                break
            spikes.append(wm)
            clean = [r for r in clean if r[0] != wm]
        spikes.sort()
        kwh = sum(vt + nt for m, vt, nt in rows)
        vt = sum(r[1] for r in rows)
        exp_vt = sum(L * T[m]["dvt"] for m, *_ in rows) / max(1e-9, sum(L * (T[m]["dvt"] + T[m]["dnt"]) for m, *_ in rows))
        kwh_spike = sum(max(0.0, (a + b) - pred(m)) for m, a, b in rows if m in spikes)
        # Вишокот во неправилните месеци: ако е распределен на ВТ/НТ како светилките, тоа е читање
        # (проценето или дополнително); ако има многу повеќе ВТ, тоа е вистински товар и преку ден.
        ex_vt = sum(a - L * T[m]["dvt"] - X * T[m]["tvt"] for m, a, b in rows if m in spikes)
        ex_tot = sum(a + b - pred(m) for m, a, b in rows if m in spikes)
        light_vt = sum(T[m]["dvt"] for m in spikes) / max(1e-9, sum(T[m]["dvt"] + T[m]["dnt"] for m in spikes))
        if not spikes or ex_tot <= 0:
            kind = "намалени читања" if spikes else ""
        elif ex_vt / ex_tot > light_vt + 0.12:
            kind = f"дополнителен товар и преку ден (ВТ {ex_vt / ex_tot:.0%} од вишокот)"
        else:
            kind = f"читање: вишокот е распределен како светилките (ВТ {ex_vt / ex_tot:.0%})"
        kwh_x = X * sum(T[m]["tvt"] + T[m]["tnt"] for m, *_ in rows if m not in spikes)
        # Сезонски профил: kWh по ноќен час, зима (ное–фев) наспроти лето (мај–авг); светилки ≈ 1.
        per_h = {m: (a + b) / (T[m]["dvt"] + T[m]["dnt"]) for m, a, b in rows if m not in spikes}
        win = [per_h[m] for m in per_h if m[5:] in ("11", "12", "01", "02")]
        sumr = [per_h[m] for m in per_h if m[5:] in ("05", "06", "07", "08")]
        season = statistics.median(win) / statistics.median(sumr) if win and sumr and statistics.median(sumr) > 0 else None
        mp = mapping.get(meter, {})
        meters.append({"meter": meter, "adresa": mp.get("address", ""), "mesto": mp.get("unit", ""), "opstina": mp.get("group", ""),
                       "meseci": len(rows), "kwh": round(kwh), "denari": round(sum(v[m][2] for m in months if m in v)),
                       "udel_vt": round(vt / kwh, 3), "udel_vt_ocekuvan": round(exp_vt, 3),
                       "nokna_kw": round(L, 3), "denonokjna_kw": round(X, 3), "kwh_denonokjno": round(kwh_x), "udel_denonokjno": round(kwh_x / kwh, 3),
                       "neprav_meseci": ", ".join(A.SHORT[m[5:]] + " " + m[2:4] for m in spikes), "kwh_neprav": round(kwh_spike),
                       "vid_nepravilnost": kind, "zima_leto_po_cas": None if season is None else round(season, 2), "ukrasuvanje": meter in deco})
    return months, hy, meters


def places(meters, hy):
    annual = V.analyse()[3]
    share = {r["opstina"]: float(r["udel_stari"]) for r in A.read("procenki.csv")}
    out = []
    for r in annual:
        ms = [m for m in meters if m["mesto"] == r["mesto"]]
        if not ms or r["mesto"] == "Кичево (град)":
            continue
        s = share[r["opstina"]]
        old = r["led"] * s / (1 - s) if r["led"] else None
        exp = None if old is None else (r["led_kw"] + old * W_EFF) * hy / (1 - V.LOSS)
        kwh = sum(m["kwh"] for m in ms)
        den = sum(m["denari"] for m in ms)
        x = sum(m["kwh_denonokjno"] for m in ms)
        sp = sum(m["kwh_neprav"] for m in ms)
        excess = None if exp is None else kwh - exp
        night = None if excess is None else excess - x - sp
        out.append({"opstina": r["opstina"], "mesto": r["mesto"], "broila": len(ms), "led": r["led"], "stari_procenka": None if old is None else round(old),
                    "kwh": kwh, "kwh_ocekuvano": None if exp is None else round(exp), "otstapuvanje_kwh": None if exp is None else round(excess),
                    "odnos": None if exp is None else round(kwh / exp, 2), "otstapuvanje_denari": None if exp is None else round(excess * den / kwh),
                    "kwh_denonokjno": x, "kwh_neprav": sp, "kwh_nokjen_visok": None if night is None else round(night),
                    "nokna_kw": round(sum(m["nokna_kw"] for m in ms), 2),
                    "nokna_kw_ocekuvano": None if old is None else round((r["led_kw"] + old * W_EFF) / (1 - V.LOSS), 2),
                    "stari_ekvivalent": None if night is None else round(max(0.0, night) * (1 - V.LOSS) / hy / W_EFF)})
    return out


def cause(p, ms):
    """Можни објаснувања за местото, по редослед на тежина."""
    c = []
    ex = max(1, p["otstapuvanje_kwh"] or 0)
    fmt = lambda v: f"{v:,}".replace(",", ".")
    if p["kwh_denonokjno"] > 0.1 * ex:
        who = ", ".join(m["meter"] for m in ms if m["kwh_denonokjno"] > 0.1 * ex)
        c.append(f"потрошувачка и преку ден {fmt(p['kwh_denonokjno'])} kWh (броило {who}): приклучен објект, кражба, светилки што светат и дење или уред во ормарот")
    if p["kwh_neprav"] > 0.1 * ex:
        when = "; ".join(f"{m['meter']}: {m['neprav_meseci']} – {m['vid_nepravilnost']}" for m in ms if m["neprav_meseci"])
        c.append(f"неправилни месеци {fmt(p['kwh_neprav'])} kWh ({when})")
    if (p["kwh_nokjen_visok"] or 0) > 0.1 * ex:
        c.append(f"ноќна потрошувачка повисока за {fmt(p['kwh_nokjen_visok'])} kWh (≈ {p['stari_ekvivalent']} стари светилки повеќе од проценката): "
                 "посилни или повеќе LED од пописот, повеќе стари светилки, осветлување на објекти или улици од друго место на истото броило")
    return "; ".join(c)


def write(meters, pl):
    by_place = defaultdict(list)
    for m in meters:
        by_place[m["mesto"]].append(m)
    for p in pl:
        p["mozni_pricini"] = cause(p, by_place[p["mesto"]])
    pl.sort(key=lambda p: -(p["otstapuvanje_kwh"] or -1e9))
    meters.sort(key=lambda m: -m["kwh_denonokjno"])
    for name, data in (("dijagnoza_broila.csv", meters), ("dijagnoza_mesta.csv", pl)):
        with open(A.IZLEZ / name, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(data[0].keys()))
            w.writeheader(); w.writerows(data)

    bold, head = Font(bold=True), PatternFill("solid", fgColor="E8EEF2")
    wb = Workbook()
    ws = wb.active
    ws.title = "Пристап"
    for t in [
        "Дијагноза на сметките за улично осветлување",
        "",
        "Фактурите ја делат енергијата на висока тарифа (ВТ, 07–22 ч, пон.–саб.) и ниска тарифа (НТ). Светилка под астрономски часовник",
        "троши во ВТ само темниот дел пред 22 ч и по 07 ч, па уделот на ВТ е однапред познат (лето околу 18 %, зима околу 34 %).",
        "Проверено на Староец и на броилото 15737 во Видрани: измереното се совпаѓа со пресметаното во рамки од 1–3 процентни поени.",
        "Ако броилото троши и преку ден (приклучок, кражба, светилки што светат и дење, уред во ормар), уделот на ВТ е поголем.",
        "",
        "За секое броило од 24 податоци (12 месеци × 2 тарифи) се пресметуваат ноќна моќност (светилки) и моќност што работи 24 часа.",
        f"Месеци што отстапуваат над {SPIKE:.0%} од пресметаното = неправилни читања (проценето читање, дополнителна наплата, корекција).",
        "Очекувано по место = (LED од пописот + стари по вашата проценка) × ноќни часови ÷ 0,9 (10 W × 80 % по стара светилка).",
        "Напомена: јуни 2026 кај сите броила има околу 30 % повеќе kWh и поголем удел на ВТ – веројатно различен период на читање.",
    ]:
        ws.append([t])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 140

    def sheet(title, header, keys, data, widths, fill_col=None):
        s = wb.create_sheet(title)
        s.append(header)
        for c in s[1]:
            c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="top")
        for r in data:
            s.append([r[k] for k in keys])
        for i, wd in enumerate(widths, 1):
            s.column_dimensions[get_column_letter(i)].width = wd
        s.freeze_panes = "C2"
        return s

    sheet("Места по отстапување", ["Поранешна општина", "Место", "Броила", "LED", "Стари (проценка)", "kWh", "Очекувано kWh", "Отстапување kWh", "Однос",
                                   "Отстапување денари", "од тоа: преку ден kWh", "од тоа: неправилни читања kWh", "од тоа: ноќно kWh", "Ноќно ≈ стари светилки", "Ноќна kW", "Очекувана ноќна kW", "Можни причини"],
          ["opstina", "mesto", "broila", "led", "stari_procenka", "kwh", "kwh_ocekuvano", "otstapuvanje_kwh", "odnos", "otstapuvanje_denari", "kwh_denonokjno",
           "kwh_neprav", "kwh_nokjen_visok", "stari_ekvivalent", "nokna_kw", "nokna_kw_ocekuvano", "mozni_pricini"], pl, [12, 20, 7, 7, 9, 9, 10, 10, 7, 11, 9, 10, 10, 9, 8, 9, 110])
    s = sheet("Броила", ["Мерно место", "Адреса", "Место", "Поранешна општина", "Месеци", "kWh", "Денари", "Удел ВТ", "Очекуван удел ВТ", "Ноќна kW", "kW 24 часа",
                         "kWh 24 часа", "Удел 24 часа", "Неправилни месеци", "kWh неправилни", "Вид на неправилност", "Зима ÷ лето по ноќен час", "Украсување"],
              ["meter", "adresa", "mesto", "opstina", "meseci", "kwh", "denari", "udel_vt", "udel_vt_ocekuvan", "nokna_kw", "denonokjna_kw", "kwh_denonokjno",
               "udel_denonokjno", "neprav_meseci", "kwh_neprav", "vid_nepravilnost", "zima_leto_po_cas", "ukrasuvanje"], meters, [16, 34, 18, 12, 7, 9, 10, 8, 9, 8, 8, 9, 8, 26, 9, 40, 9, 9])
    for row in s.iter_rows(min_row=2):
        u = row[12].value
        row[12].fill = BAD if u > 0.25 else WARN if u > 0.10 else OK
    wb.save(A.IZLEZ / "Дијагноза_сметки.xlsx")
    return pl, meters


def main():
    months, hy, meters = analyse()
    pl, meters = write(meters, places(meters, hy))
    for p in pl[:12]:
        print(p["mesto"], p["led"], p["stari_procenka"], p["kwh"], p["kwh_ocekuvano"], p["otstapuvanje_kwh"], p["odnos"], p["otstapuvanje_denari"], p["kwh_denonokjno"], p["kwh_neprav"], p["nokna_kw"], p["nokna_kw_ocekuvano"])
        print("    ", p["mozni_pricini"])
    print("--- броила со најмногу 24-часовна потрошувачка")
    for m in meters[:15]:
        print(m["meter"], m["adresa"], m["mesto"], m["kwh"], m["udel_vt"], m["udel_vt_ocekuvan"], m["nokna_kw"], m["denonokjna_kw"], m["kwh_denonokjno"], m["neprav_meseci"])
    print("вкупно 24 ч kWh:", sum(m["kwh_denonokjno"] for m in meters), "од", sum(m["kwh"] for m in meters), " неправилни:", sum(m["kwh_neprav"] for m in meters))


if __name__ == "__main__":
    main()
