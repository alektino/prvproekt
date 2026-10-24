"""Проба: колкава моќност треба да имаат старите светилки?

Бројот на стари светилки е фиксиран на вашата проценка (по поранешна општина,
и вкупно 1.500), а се пресметува колкава моќност треба да има секоја стара
светилка за да ја објасни енергијата што не ја трошат LED светилките од ЈПП:

  моќност (W) = ((kWh − украсување) × (1 − загуби) − LED kW × часови) × 1000 ÷ (број стари × удел активни × часови)

Часови: точно ноќта (зајдисонце–изгрејсонце) или ноќта + 2 часа на ден.

Употреба (по analiza.py):  python3 analiza-svetilki/alatki/proba_moknost.py
Излез: izlez/proba_moknost.csv, izlez/proba_moknost_mesecno.csv и лист „Проба – моќност“
во izlez/Анализа_светилки.xlsx
"""
import csv

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import analiza as A

LOSSES = [0.0, 0.05, 0.10]
ACTIVE = [0.8, 0.9, 1.0]
MAIN_LOSS = 0.10


def main():
    a, months, rows, by_group, compare, profiles = A.analyse()
    g_row = {(r["opstina"], r["mesec"]): r for r in by_group}
    targets = [(g, g_row[g, "all"]["stari_procenka"]) for g in A.GROUPS + ["Вкупно"]]
    targets.append(("Вкупно (1.500)", float(a["stari_procenka_vkupno"])))

    def watt(r, loss, act, n, extra):
        h = r["nokni_casovi"] + extra * r["denovi"]
        outside = (r["kwh"] - r["kwh_ukras"]) * (1 - loss) - r["led_kw"] * h
        return outside, outside * 1000 / (n * act * h) if n else None

    grid = []
    for name, n in targets:
        r = g_row["Вкупно" if name.startswith("Вкупно") else name, "all"]
        for extra in (0, A.EXTRA_MAX):
            for loss in LOSSES:
                outside, _ = watt(r, loss, 1, n, extra)
                row = {"opstina": name, "casovi": "ноќ" if not extra else f"ноќ + {extra} ч", "zagubi": loss, "led": r["led"], "stari": round(n),
                       "kwh_nadvor_od_led": round(outside)}
                for act in ACTIVE:
                    row[f"w_aktivni_{int(act * 100)}"] = round(watt(r, loss, act, n, extra)[1], 1)
                grid.append(row)

    total_n = float(a["stari_procenka_vkupno"])
    monthly = []
    for m in months:
        r = g_row["Вкупно", m]
        row = {"mesec": m, "nokni_casovi": r["nokni_casovi"], "kwh": round(r["kwh"]), "kwh_ukras": round(r["kwh_ukras"]), "kwh_led_nok": round(r["kwh_led_nok"])}
        for extra in (0, A.EXTRA_MAX):
            for act in (0.8, 1.0):
                row[f"w_{'nok' if not extra else 'nok2'}_{int(act * 100)}"] = round(watt(r, MAIN_LOSS, act, total_n, extra)[1], 1)
        monthly.append(row)

    for name, data in (("proba_moknost.csv", grid), ("proba_moknost_mesecno.csv", monthly)):
        with open(A.IZLEZ / name, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(data[0].keys()))
            w.writeheader(); w.writerows(data)
    write_sheet(grid, monthly, total_n)
    for r in grid:
        if r["opstina"].startswith("Вкупно") and r["zagubi"] == MAIN_LOSS:
            print(f"{r['opstina']:15} {r['casovi']:9} стари {r['stari']:5}: надвор од LED {r['kwh_nadvor_od_led']:8,} kWh → W при 80/90/100 %: "
                  f"{r['w_aktivni_80']} / {r['w_aktivni_90']} / {r['w_aktivni_100']}")


def write_sheet(grid, monthly, total_n):
    path = A.IZLEZ / "Анализа_светилки.xlsx"
    wb = load_workbook(path)
    if "Проба – моќност" in wb.sheetnames:
        del wb["Проба – моќност"]
    ws = wb.create_sheet("Проба – моќност", 1)
    bold, head, hl = Font(bold=True), PatternFill("solid", fgColor="E8EEF2"), PatternFill("solid", fgColor="E3F1E3")
    ws.append(["Проба: колкава моќност треба да имаат старите светилки за да ја објаснат енергијата надвор од LED?"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Моќност (W) = ((kWh − украсување) × (1 − загуби) − LED kW × часови) × 1000 ÷ (број стари × удел активни × часови)."])
    ws.append(["Стари = вашата проценка по поранешна општина (збир 1.731), и посебно вкупно 1.500. Часови = точно ноќта или ноќта + 2 часа на ден."])
    ws.append([])
    ws.append(["Поранешна општина", "Часови", "Загуби", "LED", "Стари", "kWh надвор од LED", "W при 80 % активни", "W при 90 %", "W при 100 %"])
    for c in ws[ws.max_row]:
        c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="top")
    for r in grid:
        ws.append([r["opstina"], r["casovi"], r["zagubi"], r["led"], r["stari"], r["kwh_nadvor_od_led"], r["w_aktivni_80"], r["w_aktivni_90"], r["w_aktivni_100"]])
        row = ws[ws.max_row]
        row[2].number_format = "0%"
        for c in row[3:6]:
            c.number_format = "#,##0"
        if r["opstina"].startswith("Вкупно"):
            for c in row:
                c.font = bold
            if r["zagubi"] == MAIN_LOSS:
                for c in row:
                    c.fill = hl
    ws.append([])
    ws.append([f"По месец (сите општини, загуби {MAIN_LOSS:.0%}, {total_n:,.0f} стари светилки)".replace(",", ".")])
    ws[ws.max_row][0].font = bold
    ws.append(["Месец", "Ноќни часови", "kWh", "Украсување", "LED цела ноќ kWh", "W (ноќ, 80 %)", "W (ноќ, 100 %)", "W (ноќ + 2 ч, 80 %)", "W (ноќ + 2 ч, 100 %)"])
    for c in ws[ws.max_row]:
        c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="top")
    for r in monthly:
        ws.append([r["mesec"], r["nokni_casovi"], r["kwh"], r["kwh_ukras"], r["kwh_led_nok"], r["w_nok_80"], r["w_nok_100"], r["w_nok2_80"], r["w_nok2_100"]])
        for c in ws[ws.max_row][2:5]:
            c.number_format = "#,##0"
    for col, w in zip("ABCDEFGHI", [18, 10, 9, 8, 8, 14, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    wb.save(path)


if __name__ == "__main__":
    main()
